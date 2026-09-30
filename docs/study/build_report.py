"""Export the reviewed Markdown report to offline HTML and Word with diagrams."""
import base64
from pathlib import Path

import markdown
from lxml import html
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
source = (HERE / 'REPORT_RU.md').read_text(encoding='utf-8')
body = markdown.markdown(source, extensions=['tables', 'fenced_code'])
offline = body
for name in ('database', 'question_crud', 'search'):
    data = base64.b64encode((HERE / f'{name}.svg').read_bytes()).decode('ascii')
    offline = offline.replace(f'src="{name}.svg"', f'src="data:image/svg+xml;base64,{data}"')
css = '''
body{font:16px/1.65 "Segoe UI",Arial,sans-serif;color:#17283a;background:#edf2f7;margin:0}
main{max-width:1100px;margin:30px auto;padding:48px;background:white;border-radius:12px}
h1{font-size:34px;line-height:1.2;color:#123d65}h2{margin-top:45px;color:#123d65;border-top:2px solid #d5e5f2;padding-top:24px}
h3{color:#245b7f}table{border-collapse:collapse;width:100%;font-size:14px;margin:18px 0}
td,th{border:1px solid #cad6e2;padding:9px 12px;text-align:left;vertical-align:top}th{background:#e8f1f8}
tr:nth-child(even){background:#f7f9fc}a{color:#0969a2}img{width:100%;height:auto;max-height:none}
code{font-family:Consolas,monospace;font-size:.91em;overflow-wrap:anywhere}pre{background:#f2f5f8;padding:18px;border-radius:6px;white-space:pre-wrap}
@media print{body{background:white;font-size:10pt}main{margin:0;padding:0;max-width:none}h2,h3{break-after:avoid}tr,img,pre{break-inside:avoid}a{color:inherit}img{max-height:240mm;object-fit:contain}}
@media(max-width:700px){main{margin:0;padding:20px}table{font-size:12px}td,th{padding:5px}h1{font-size:26px}}
'''
(HERE / 'REPORT_RU.html').write_text('<!doctype html><html lang="ru"><meta charset="utf-8">'
    '<meta name="viewport" content="width=device-width, initial-scale=1"><title>STACK UNDERFLOW — учебный отчёт</title>'
    f'<style>{css}</style><main>{offline}</main></html>', encoding='utf-8')

document = Document()
document.core_properties.title = 'STACK UNDERFLOW: MVP, база данных, CRUD и поиск'
document.core_properties.subject = 'Проверка текущей реализации для учебной защиты'
document.core_properties.author = 'STACK UNDERFLOW project'
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(7)
for style in ('Heading 1', 'Heading 2', 'Heading 3'):
    document.styles[style].font.color.rgb = RGBColor.from_string('123D65')


def layout(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11.69 if landscape else 8.27)
    section.page_height = Inches(8.27 if landscape else 11.69)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.65)


layout(document.sections[0])
footer = document.sections[0].footer.paragraphs[0]
footer.add_run('STACK UNDERFLOW | Учебный отчёт | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def inline(paragraph, element, bold=False, italic=False, code=False):
    def text(value, b=bold, i=italic, c=code):
        if not value:
            return
        run = paragraph.add_run(value)
        run.bold, run.italic = b, i
        if c:
            run.font.name = 'Consolas'
            run.font.size = Pt(8)
    text(element.text)
    for child in element:
        if child.tag == 'br':
            paragraph.add_run().add_break()
        else:
            inline(paragraph, child, bold or child.tag == 'strong',
                   italic or child.tag == 'em', code or child.tag == 'code')
            if child.tag == 'a':
                # Keep relative artifact/code references readable in a printed report.
                text(f" ({child.get('href')})", c=True)
        text(child.tail)


for element in html.fragment_fromstring(body, create_parent=True):
    if element.tag in ('h1', 'h2', 'h3'):
        document.add_heading(element.text_content(), level=int(element.tag[1]) - 1)
    elif element.tag == 'p' and element.find('img') is not None:
        image = element.find('img')
        path = HERE / Path(image.get('src')).with_suffix('.png')
        width, height = Image.open(path).size
        landscape = width > height * 1.15
        layout(document.add_section(WD_SECTION.NEW_PAGE), landscape)
        document.add_heading(image.get('alt'), level=2)
        max_width, max_height = (10.3, 6.15) if landscape else (6.9, 9.5)
        scale = min(max_width / width, max_height / height)
        document.add_picture(str(path), width=Inches(width * scale), height=Inches(height * scale))
        layout(document.add_section(WD_SECTION.NEW_PAGE))
    elif element.tag == 'p':
        inline(document.add_paragraph(), element)
    elif element.tag == 'pre':
        paragraph = document.add_paragraph()
        run = paragraph.add_run(element.text_content().rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
    elif element.tag == 'table':
        rows = element.xpath('.//tr')
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for row_index, row in enumerate(rows):
            cells = table.add_row().cells
            for index, cell in enumerate(row):
                inline(cells[index].paragraphs[0], cell, bold=row_index == 0)
            if row_index == 0:
                repeat = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
        document.add_paragraph()
    elif element.tag in ('ul', 'ol'):
        for item in element:
            inline(document.add_paragraph(style='List Bullet' if element.tag == 'ul' else 'List Number'), item)
document.save(HERE / 'REPORT_RU.docx')
print('Created REPORT_RU.html and REPORT_RU.docx')
